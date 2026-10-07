"""Daily log v9 (lot 10) sur un VRAI PostgreSQL : routes JSON, repository,
migration 0014 (triggers), verrou de version, journal d'audit.

Chaque règle des décisions du gérant y est jouée de bout en bout : la
requête HTTP passe par la vraie validation ET par les vrais triggers.
"""
import datetime
import threading

from tests_integration.base import BaseIntegration


class _BaseDL(BaseIntegration):
    def setUp(self):
        super().setUp()
        from app.repositories import dailylog
        self.dl = dailylog
        self.chef = self.user("chef_de_projet")
        self.moi = self.user("intervenant")
        self.p1 = self.projet(self.chef, code="26069D")
        self.p2 = self.projet(self.chef, code="26006P")
        self.sql("UPDATE utilisateur SET created_at = now() - interval '400 days'")
        self.auj = self.dl.aujourdhui_pour({"fuseau": "Africa/Tunis"})

    def jour(self, n):
        return self.auj - datetime.timedelta(days=n)

    def put(self, date, lignes, version=None, qs="", **extra):
        if version is None:
            version = self.version(self.moi_courant, date)
        corps = {"version": version, "lignes": lignes, **extra}
        return self.client.put(f"/dailylog/api/jour/{date.isoformat()}{qs}", json=corps)

    def version(self, uid, date):
        r = self.sql("SELECT version FROM dailylog_jour WHERE utilisateur_id=%s AND date=%s", (uid, date))
        return r[0]["version"] if r else 0

    def lignes(self, uid, date):
        return [(r["projet_id"], r["tache_id"], float(r["heures"])) for r in self.sql(
            "SELECT projet_id, tache_id, heures FROM dailylog_entree WHERE utilisateur_id=%s AND date=%s ORDER BY projet_id, tache_id",
            (uid, date))]

    def connecter(self, uid):
        self.moi_courant = uid
        self.login(uid)

    def historique(self, uid, date, lignes, absent=None):
        """Données anciennes (v2, Chronos) écrites comme la reprise : hors règles."""
        with self.db.get_cursor() as cur:
            cur.execute("SET LOCAL kairos.reprise_historique = 'on'")
            for projet_id, heures in lignes:
                cur.execute("INSERT INTO dailylog_entree (utilisateur_id, date, projet_id, heures) VALUES (%s,%s,%s,%s)",
                            (uid, date, projet_id, heures))
            if absent is not None:
                cur.execute("INSERT INTO dailylog_jour (utilisateur_id, date, absent) VALUES (%s,%s,%s)", (uid, date, absent))


class TestSaisie(_BaseDL):
    def test_enregistrer_relire_version_et_renvoi_identique_sans_ecriture(self):
        self.connecter(self.moi)
        r = self.put(self.auj, [{"projet_id": self.p1, "tache_id": None, "minutes": 270}])
        self.assertEqual(r.status_code, 200, r.data[:400])
        corps = r.get_json()
        self.assertEqual((corps["version"], corps["statut"], corps["total_minutes"], corps["enregistre"]), (1, "rempli", 270, True))
        self.assertEqual(self.lignes(self.moi, self.auj), [(self.p1, None, 4.5)])
        audit = self.sql("SELECT count(*) AS n FROM audit_log")[0]["n"]
        r = self.put(self.auj, [{"projet_id": self.p1, "minutes": 270}], version=1)
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.get_json()["enregistre"])
        self.assertEqual(self.version(self.moi, self.auj), 1)
        self.assertEqual(self.sql("SELECT count(*) AS n FROM audit_log")[0]["n"], audit)
        # Diff : seule la ligne changée est réécrite.
        r = self.put(self.auj, [{"projet_id": self.p1, "minutes": 270}, {"projet_id": self.p2, "minutes": 60}], version=1)
        self.assertEqual(r.get_json()["version"], 2)
        self.assertEqual(self.sql("SELECT count(*) AS n FROM audit_log WHERE table_cible='dailylog_entree'")[0]["n"], 2)
        g = self.client.get(f"/dailylog/api/jour/{self.auj}").get_json()
        self.assertEqual([(l["projet_id"], l["minutes"]) for l in g["lignes"]], [(self.p1, 270), (self.p2, 60)])

    def test_version_perimee_409(self):
        self.connecter(self.moi)
        self.put(self.auj, [{"projet_id": self.p1, "minutes": 60}], version=0)
        r = self.put(self.auj, [{"projet_id": self.p1, "minutes": 120}], version=0)
        self.assertEqual(r.status_code, 409)
        self.assertEqual(r.get_json()["jour"]["version"], 1)
        self.assertEqual(self.lignes(self.moi, self.auj), [(self.p1, None, 1.0)])

    def test_deux_onglets_simultanes_un_200_un_409(self):
        self.connecter(self.moi)
        resultats = []
        depart = threading.Barrier(2)

        def onglet(minutes):
            client = self.app.test_client()
            from app.auth import password_fingerprint
            with client.session_transaction() as s:
                s["user_id"] = self.moi
                s["pw_fingerprint"] = password_fingerprint("h")
            depart.wait()
            r = client.put(f"/dailylog/api/jour/{self.auj}",
                           json={"version": 0, "lignes": [{"projet_id": self.p1, "minutes": minutes}]})
            resultats.append(r.status_code)

        fils = [threading.Thread(target=onglet, args=(m,)) for m in (240, 300)]
        for f in fils:
            f.start()
        for f in fils:
            f.join()
        self.assertEqual(sorted(resultats), [200, 409])
        self.assertEqual(len(self.lignes(self.moi, self.auj)), 1)
        self.assertEqual(self.version(self.moi, self.auj), 1)

    def test_pas_de_30_min_et_10_h(self):
        self.connecter(self.moi)
        r = self.put(self.auj, [{"projet_id": self.p1, "minutes": 135}])
        self.assertEqual([e["code"] for e in r.get_json()["erreurs"]], ["pas_30_min"])
        r = self.put(self.auj, [{"projet_id": self.p1, "minutes": 600}])
        self.assertEqual(r.status_code, 200)
        r = self.put(self.auj, [{"projet_id": self.p1, "minutes": 600}, {"projet_id": self.p2, "minutes": 30}])
        self.assertIn("max_journee", [e["code"] for e in r.get_json()["erreurs"]])
        self.assertEqual(self.lignes(self.moi, self.auj), [(self.p1, None, 10.0)])

    def test_donnees_anciennes_relues_reenregistrees_reduites(self):
        d = self.jour(3)
        self.historique(self.moi, d, [(self.p1, 2.25), (self.p2, 9.75)])  # 12 h, quart d'heure
        self.connecter(self.moi)
        j = self.client.get(f"/dailylog/api/jour/{d}").get_json()
        self.assertEqual([(l["minutes"], l["hors_grille"]) for l in j["lignes"]], [(135, True), (585, True)])
        identique = [{"projet_id": self.p1, "minutes": 135}, {"projet_id": self.p2, "minutes": 585}]
        r = self.put(d, identique)
        self.assertEqual(r.status_code, 200, r.data[:400])
        self.assertFalse(r.get_json()["enregistre"])
        # réduire une ligne d'une journée de 12 h : accepté (et la ligne au quart d'heure reste telle quelle)
        r = self.put(d, [{"projet_id": self.p1, "minutes": 135}, {"projet_id": self.p2, "minutes": 540}])
        self.assertEqual(r.status_code, 200, r.data[:400])
        self.assertEqual(self.lignes(self.moi, d), [(self.p1, None, 2.25), (self.p2, None, 9.0)])
        # l'augmenter : refusé (10 h)
        r = self.put(d, [{"projet_id": self.p1, "minutes": 150}, {"projet_id": self.p2, "minutes": 540}])
        self.assertIn("max_journee", [e["code"] for e in r.get_json()["erreurs"]])

    def test_vider_seulement_avec_confirmation(self):
        self.connecter(self.moi)
        self.put(self.auj, [{"projet_id": self.p1, "minutes": 240}])
        r = self.put(self.auj, [])
        self.assertEqual([e["code"] for e in r.get_json()["erreurs"]], ["journee_vide"])
        r = self.put(self.auj, [], vider=True)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()["statut"], "vide")
        self.assertEqual(self.lignes(self.moi, self.auj), [])

    def test_absence_et_annulation_qui_rend_les_lignes(self):
        self.connecter(self.moi)
        self.put(self.auj, [{"projet_id": self.p1, "minutes": 240}])
        url = f"/dailylog/api/jour/{self.auj}/absence"
        r = self.client.post(url, json={"version": 1, "absent": True})
        self.assertEqual(r.status_code, 200, r.data[:400])
        self.assertEqual(r.get_json()["statut"], "absent")
        self.assertEqual(self.lignes(self.moi, self.auj), [])
        r = self.client.post(url, json={"version": 2, "absent": False, "lignes": [{"projet_id": self.p1, "minutes": 240}]})
        self.assertEqual(r.get_json()["statut"], "rempli")
        self.assertEqual(self.lignes(self.moi, self.auj), [(self.p1, None, 4.0)])
        r = self.client.post(url, json={"version": 3, "absent": True, "lignes": [{"projet_id": self.p1, "minutes": 60}]})
        self.assertIn("absent_avec_heures", [e["code"] for e in r.get_json()["erreurs"]])
        # Un jour absent n'est pas « manqué ».
        self.assertTrue(self.dl.jour_renseigne(self.moi, self.auj))

    def test_projets_bloque_clos_pas_commence_tache(self):
        bloque = self.projet(self.chef, etat="bloque")
        futur = self.projet(self.chef, debut=self.auj + datetime.timedelta(days=5))
        clos = self.projet(self.chef)
        self.sql("UPDATE projet SET etat='termine', date_cloture=%s WHERE id=%s", (self.jour(2), clos))
        autre_tache = self.tache(self.p2, self.chef)
        self.connecter(self.moi)
        self.assertEqual(self.put(self.auj, [{"projet_id": bloque, "minutes": 60}]).status_code, 200)
        self.assertEqual(self.put(self.jour(2), [{"projet_id": clos, "minutes": 60}]).status_code, 200)  # jour de la clôture
        codes = lambda r: [e["code"] for e in r.get_json()["erreurs"]]
        self.assertEqual(codes(self.put(self.jour(1), [{"projet_id": clos, "minutes": 60}])), ["projet_clos"])
        self.assertEqual(codes(self.put(self.jour(1), [{"projet_id": futur, "minutes": 60}])), ["projet_pas_commence"])
        self.assertEqual(codes(self.put(self.jour(1), [{"projet_id": self.p1, "tache_id": autre_tache, "minutes": 60}])),
                         ["tache_hors_projet"])
        ma_tache = self.tache(self.p1, self.chef)
        r = self.put(self.jour(1), [{"projet_id": self.p1, "tache_id": ma_tache, "minutes": 60},
                                    {"projet_id": self.p1, "minutes": 30}])
        self.assertEqual(r.status_code, 200, r.data[:400])
        self.assertEqual(self.lignes(self.moi, self.jour(1)), [(self.p1, ma_tache, 1.0), (self.p1, None, 0.5)])

    def test_projet_d_une_autre_equipe_invisible(self):
        chef_urbs = self.user("chef_de_projet", "URBS")
        urbs = self.projet(chef_urbs, "URBS")
        self.connecter(self.moi)
        r = self.put(self.auj, [{"projet_id": urbs, "minutes": 60}])
        self.assertEqual([e["code"] for e in r.get_json()["erreurs"]], ["projet_invisible"])
        r = self.put(self.auj, [{"projet_id": 999999, "minutes": 60}])
        self.assertEqual([e["code"] for e in r.get_json()["erreurs"]], ["projet_invisible"])


class TestFenetreEtCorrection(_BaseDL):
    def test_fenetre_de_14_jours(self):
        self.connecter(self.moi)
        self.assertEqual(self.put(self.jour(14), [{"projet_id": self.p1, "minutes": 480}]).status_code, 200)
        r = self.put(self.jour(15), [{"projet_id": self.p1, "minutes": 480}])
        self.assertEqual([e["code"] for e in r.get_json()["erreurs"]], ["ligne_non_modifiable"])
        self.assertEqual(self.lignes(self.moi, self.jour(15)), [])
        j = self.client.get(f"/dailylog/api/jour/{self.jour(15)}").get_json()
        self.assertTrue(j["droits"]["lecture_seule"])
        # Garde-fou en base, même sans passer par la route.
        with self.assertRaises(Exception) as ctx:
            with self.db.get_cursor(user_id=self.moi) as cur:
                cur.execute("INSERT INTO dailylog_entree (utilisateur_id, date, projet_id, heures) VALUES (%s,%s,%s,1)",
                            (self.moi, self.jour(15), self.p1))
        self.assertEqual(getattr(ctx.exception, "pgcode", None), "42501")

    def test_correction_par_le_chef_tracee_a_son_nom(self):
        chef2 = self.user("chef_de_projet")
        p_autre = self.projet(chef2)
        d = self.jour(40)
        self.historique(self.moi, d, [(self.p1, 7), (p_autre, 5)])  # 12 h (Chronos)
        self.connecter(self.chef)
        qs = f"?utilisateur={self.moi}"
        j = self.client.get(f"/dailylog/api/jour/{d}{qs}").get_json()
        self.assertEqual({l["projet_id"]: l["modifiable"] for l in j["lignes"]}, {self.p1: True, p_autre: False})
        dernier_audit = self.sql("SELECT coalesce(max(id), 0) AS m FROM audit_log")[0]["m"]
        # sa ligne ramenée à 5 h ; celle de l'autre chef, absente de la requête, est gardée
        r = self.client.put(f"/dailylog/api/jour/{d}{qs}", json={"version": 0, "lignes": [{"projet_id": self.p1, "minutes": 300}]})
        self.assertEqual(r.status_code, 200, r.data[:400])
        self.assertEqual(self.lignes(self.moi, d), [(self.p1, None, 5.0), (p_autre, None, 5.0)])
        maj = self.sql("SELECT updated_by FROM dailylog_entree WHERE utilisateur_id=%s AND date=%s AND projet_id=%s",
                       (self.moi, d, self.p1))[0]["updated_by"]
        self.assertEqual(maj, self.chef)
        audit = self.sql("SELECT table_cible, utilisateur_id FROM audit_log WHERE id > %s ORDER BY id", (dernier_audit,))
        self.assertEqual([(a["table_cible"], a["utilisateur_id"]) for a in audit],
                         [("dailylog_entree", self.chef), ("dailylog_jour", self.chef)])
        notif = self.sql("SELECT message FROM notification WHERE utilisateur_id=%s", (self.moi,))
        self.assertEqual(len(notif), 1)
        self.assertIn(f"Daily log du {d:%d/%m/%Y} corrigé par", notif[0]["message"])
        # modifier la ligne de l'autre chef : refusé ; l'absence : refusée
        r = self.client.put(f"/dailylog/api/jour/{d}{qs}", json={"version": 1, "lignes": [
            {"projet_id": self.p1, "minutes": 300}, {"projet_id": p_autre, "minutes": 240}]})
        self.assertEqual([e["code"] for e in r.get_json()["erreurs"]], ["ligne_non_modifiable"])
        r = self.client.post(f"/dailylog/api/jour/{d}/absence{qs}", json={"version": 1, "absent": True})
        self.assertIn("absence_non_autorisee", [e["code"] for e in r.get_json()["erreurs"]])
        # un intervenant n'ouvre pas le Daily log d'un autre
        self.connecter(self.user("intervenant"))
        self.assertEqual(self.client.get(f"/dailylog/api/jour/{d}{qs}").status_code, 403)

    def test_ligne_masquee_pour_un_chef_d_une_autre_equipe(self):
        chef_urbs = self.user("chef_de_projet", "URBS")
        urbs = self.projet(chef_urbs, "URBS", code="URBS01X")
        self.sql("INSERT INTO projet_intervenant (projet_id, utilisateur_id) VALUES (%s, %s)", (urbs, self.moi))
        self.historique(self.moi, self.jour(1), [(self.p1, 4), (urbs, 3)])
        self.connecter(chef_urbs)
        r = self.client.get(f"/dailylog/api/jour/{self.jour(1)}?utilisateur={self.moi}")
        corps = r.get_json()
        self.assertNotIn("26069D", r.data.decode())
        self.assertEqual(corps["total_minutes"], 420)
        self.assertEqual(sorted((l["masquee"], l["minutes"]) for l in corps["lignes"]), [(False, 180), (True, 240)])
        # Corriger sa ligne garde la ligne masquée.
        r = self.client.put(f"/dailylog/api/jour/{self.jour(1)}?utilisateur={self.moi}",
                            json={"version": 0, "lignes": [{"projet_id": urbs, "minutes": 120}]})
        self.assertEqual(r.status_code, 200, r.data[:300])
        self.assertEqual(self.lignes(self.moi, self.jour(1)), sorted([(self.p1, None, 4.0), (urbs, None, 2.0)]))

    def test_admin_corrige_tout_absence_comprise(self):
        admin = self.user("admin", equipe=None)
        d = self.jour(90)
        self.historique(self.moi, d, [(self.p1, 8)])
        self.connecter(admin)
        r = self.client.post(f"/dailylog/api/jour/{d}/absence?utilisateur={self.moi}", json={"version": 0, "absent": True})
        self.assertEqual(r.status_code, 200, r.data[:300])
        self.assertEqual(self.lignes(self.moi, d), [])
        self.assertEqual(self.sql("SELECT updated_by FROM dailylog_jour WHERE utilisateur_id=%s AND date=%s",
                                  (self.moi, d))[0]["updated_by"], admin)

    def test_fuseau_de_la_guyane(self):
        self.sql("UPDATE utilisateur SET fuseau='America/Cayenne' WHERE id=%s", (self.moi,))
        auj_guyane = self.dl.aujourdhui_pour({"fuseau": "America/Cayenne"})
        self.connecter(self.moi)
        r = self.client.get(f"/dailylog/api/jour/{auj_guyane + datetime.timedelta(days=1)}")
        self.assertEqual(r.status_code, 422)
        self.assertEqual(r.get_json()["aujourdhui"], auj_guyane.isoformat())
        self.assertEqual(self.put(auj_guyane, [{"projet_id": self.p1, "minutes": 60}]).status_code, 200)


class TestLecturesEtRecherche(_BaseDL):
    def test_recherche_sans_accents_jokers_bloque_clos_visibilite(self):
        ete = self.projet(self.chef, code="25140P")
        self.sql("UPDATE projet SET nom='Résidence Été' WHERE id=%s", (ete,))
        bloque = self.projet(self.chef, etat="bloque", code="BLOQ1X")
        clos = self.projet(self.chef, code="CLOS1X")
        self.sql("UPDATE projet SET etat='termine', date_cloture=%s, nom='Clos 100%%' WHERE id=%s", (self.jour(5), clos))
        urbs = self.projet(self.user("chef_de_projet", "URBS"), "URBS", code="URBS2X")
        self.connecter(self.moi)
        trouve = lambda q, date=None: {r["projet_id"] for r in self.client.get(
            "/dailylog/recherche-projets", query_string={"q": q, **({"date": date.isoformat()} if date else {})}).get_json()["resultats"]}
        self.assertEqual(trouve("ete"), {ete})
        self.assertEqual(trouve("ÉTÉ"), {ete})
        self.assertIn(bloque, trouve("bloq"))
        self.assertEqual(trouve("clos"), set())                     # clos avant aujourd'hui
        self.assertEqual(trouve("clos", self.jour(6)), {clos})      # saisie antidatée
        self.assertEqual(trouve("100%", self.jour(6)), {clos})
        self.assertEqual(trouve("%%"), set())                       # « % » n'est pas un joker
        self.assertEqual(trouve("__"), set())
        self.assertEqual(trouve("urbs"), set())                     # autre équipe : invisible
        self.sql("UPDATE projet SET nom='Pro_jet' WHERE id=%s", (self.p2,))
        self.assertEqual(trouve("o_j"), {self.p2})

    def test_taches_d_un_projet(self):
        t1 = self.tache(self.p1, self.chef)
        self.tache(self.p1, self.chef, etat="termine")
        self.connecter(self.moi)
        r = self.client.get("/dailylog/api/taches", query_string={"projet": self.p1}).get_json()
        self.assertEqual([t["tache_id"] for t in r["taches"]], [t1])
        urbs = self.projet(self.user("chef_de_projet", "URBS"), "URBS")
        self.assertEqual(self.client.get("/dailylog/api/taches", query_string={"projet": urbs}).status_code, 404)

    def test_catalogue_mois_objectif_et_page(self):
        t = self.tache(self.p1, self.chef, intervenants=[self.moi])
        self.historique(self.moi, self.jour(30), [(self.p2, 8)])
        self.connecter(self.moi)
        cat = self.dl.list_lignes_suggerees(self.moi)
        self.assertEqual([(c["projet_id"], c["tache_id"]) for c in cat["mine"]], [(self.p1, None), (self.p1, t)])
        self.assertEqual([c["projet_id"] for c in cat["recentes"]], [self.p2])
        self.assertEqual(cat["mine"][0]["nb_taches"], 1)
        r = self.client.put("/dailylog/api/objectif", json={"objectif_minutes": 450})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.dl.get_objectif(self.moi), 450)
        self.assertEqual(self.client.put("/dailylog/api/objectif", json={"objectif_minutes": 275}).status_code, 422)
        self.historique(self.moi, self.jour(1), [], absent=True)
        mois = self.client.get(f"/dailylog/api/mois/{self.jour(1):%Y-%m}").get_json()["jours"]
        self.assertEqual(mois.get(self.jour(1).isoformat()), "absent")
        page = self.client.get("/dailylog")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'"objectifMinutes": 450', page.data)

    def test_jours_manques_bornes_par_l_arrivee(self):
        nouveau = self.user("intervenant")
        self.sql("UPDATE utilisateur SET date_embauche=%s WHERE id=%s", (self.jour(3), nouveau))
        u = self.sql("SELECT id, fuseau, COALESCE(date_embauche, created_at::date) AS debut_daily_log FROM utilisateur WHERE id=%s",
                     (nouveau,))[0]
        manques = self.dl.jours_manques_recents(u)
        self.assertTrue(all(j >= self.jour(3) and j.weekday() < 5 for j in manques))
        ancien = self.sql("SELECT id, fuseau, COALESCE(date_embauche, created_at::date) AS debut_daily_log FROM utilisateur WHERE id=%s",
                          (self.moi,))[0]
        self.assertGreaterEqual(len(self.dl.jours_manques_recents(ancien)), 9)

    def test_profil_une_requete_par_periode_et_absences(self):
        self.historique(self.moi, self.jour(1), [(self.p1, 3)])
        self.historique(self.moi, self.jour(2), [], absent=True)
        self.connecter(self.chef)
        html = self.client.get(f"/utilisateurs/{self.moi}/profil").data.decode()
        self.assertIn("Absent", html)
        self.assertIn(f"/dailylog?date={self.jour(1)}&amp;utilisateur={self.moi}", html)
        lignes = self.dl.list_entrees_periode(self.moi, self.jour(7), self.auj)
        self.assertEqual([(l["date"], float(l["heures"])) for l in lignes], [(self.jour(1), 3.0)])
        self.assertEqual(self.dl.jours_absents(self.moi, self.jour(7), self.auj), {self.jour(2)})

    def test_heures_chef_intervenant_et_vues_inchangees(self):
        self.connecter(self.moi)
        self.put(self.auj, [{"projet_id": self.p1, "minutes": 150}])
        self.connecter(self.chef)
        self.put(self.auj, [{"projet_id": self.p1, "minutes": 60}])
        r = self.sql("SELECT heures_chef, heures_intervenant FROM v_projet_heures_par_role WHERE projet_id=%s", (self.p1,))[0]
        self.assertEqual((float(r["heures_chef"]), float(r["heures_intervenant"])), (1.0, 2.5))
        self.assertEqual(float(self.sql("SELECT heures_cumulees FROM v_projet_heures WHERE projet_id=%s", (self.p1,))[0]["heures_cumulees"]), 3.5)

    def test_export_excel_et_liste_des_projets_lisent_les_heures(self):
        import io
        from openpyxl import load_workbook
        self.historique(self.moi, self.jour(30), [(self.p1, 2.67)])   # ancienne valeur Chronos
        self.connecter(self.moi)
        self.put(self.auj, [{"projet_id": self.p1, "minutes": 150}])
        self.connecter(self.chef)
        self.put(self.auj, [{"projet_id": self.p1, "minutes": 60}])
        r = self.client.get("/projets/export.xlsx")
        self.assertEqual(r.status_code, 200)
        feuille = load_workbook(io.BytesIO(r.data)).active
        entetes = [c.value for c in feuille[1]]
        ligne = next(row for row in feuille.iter_rows(min_row=2, values_only=True) if row[0] == "26069D")
        self.assertAlmostEqual(ligne[entetes.index("Heures intervenants")], 5.17, places=2)
        self.assertAlmostEqual(ligne[entetes.index("Heures chef de projet")], 1.0, places=2)
